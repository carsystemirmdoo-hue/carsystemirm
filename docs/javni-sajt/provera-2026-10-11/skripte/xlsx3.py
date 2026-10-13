import json, csv, re, sys, os, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
REPO="/Users/miledulic/Carsystem/radno/javni-sajt-ispravke"
OUT=sys.argv[1]
rows=json.load(open('slike.json'))
unres={}
for r in csv.DictReader(open(f"{REPO}/docs/catalog/image-gaps-2026-10/UNRESOLVED_IMAGES.csv")):
    slug=r['image_id'].split('__',1)[-1] if r['image_id'] else ''
    unres[slug]=r
sata={r['slug']:r for r in csv.DictReader(open(f"{REPO}/docs/catalog/image-gaps-2026-10/sata-current-official-urls.csv"))}
BR={"rm":"R-M","baslac":"baslac","norbin":"Norbin","cosmos-lac":"Cosmos Lac","sata":"SATA","rupes":"RUPES"}
ORDER={"rm":0,"baslac":1,"norbin":2,"cosmos-lac":3,"sata":4,"rupes":5}
SITE="https://carsystemirm.com"
INTERNAL=re.compile(r'^(RM-|NORBIN-|CL-|SATA-)')
def code_of(x):
    code,note=_code_of(x)
    if code and not re.search(r'\d', code):
        return "", "sistem bez šifre proizvođača"
    if code and x['brand']=='rupes' and re.fullmatch(r'\d{3}\.\d{3}', code):
        return "", f"Carsystem broj artikla {code}; RUPES šifra nije potvrđena"
    return code,note
def _code_of(x):
    if x['brand']=='sata' and x['slug'] in sata and sata[x['slug']]['main_article']:
        return sata[x['slug']]['main_article'], "SATA broj artikla (sa sata.com)"
    if x['manufacturerCode']: return x['manufacturerCode'], ""
    for s in x['specs']:
        lab,val=s.split(': ',1)
        if val and not INTERNAL.match(val): return val, ""
    if x['publicCode'] and not INTERNAL.match(x['publicCode']): return x['publicCode'], ""
    if x['externalSku'] and not INTERNAL.match(x['externalSku']): return x['externalSku'], ""
    return "", "šifra proizvođača nije potvrđena"
def pack_of(x):
    parts=[]
    if x['variantOptions']: parts+= [v for v in x['variantOptions'].values() if v and not re.fullmatch(r'[a-z0-9]+(-[a-z0-9]+)+', v)]
    p=[l for l in x['packages'] if l and l!='Na upit']
    parts+=[l for l in p if l not in parts]
    if x['volume'] and x['volume'] not in parts: parts.append(x['volume'])
    return ", ".join(parts)
def strip_brand(name,b):
    return re.sub(r'^(R-M|baslac|Baslac|Norbin|Cosmos Lac|SATA|RUPES)\s+','',name)
def query_of(x,code):
    b=BR[x['brand']]; n=strip_brand(x['name'],b)
    if x['brand']=='sata': return f'SATA {code} {n}'.strip()
    q=f'{b} {code} {n}' if code and code not in n else f'{b} {n}'
    return re.sub(r'\s+',' ',q).strip()+' packshot'
def source_of(x):
    if x['brand']=='sata' and x['slug'] in sata:
        s=sata[x['slug']]; return f"Stranica artikla: {s['product_page_url']} · zvanična slika: {s['current_image_url']}"
    u=unres.get(x['slug'])
    return (u['checked_sources'] if u else '')
def note_of(x):
    u=unres.get(x['slug'])
    base=u['reason'] if u else ''
    if x['brand']=='sata':
        if u and 'PLACEHOLDER_ACCEPTED' in (u['state']+u['reason']+u.get('group','')): base='Placeholder ranije prihvaćen (pribor). '+base
        pass
    if x['brand']=='cosmos-lac':
        if x['slug']!='cosmos-lac-easy-max-glitter-912-multi':
            base='ISPRAVKA 2026-10-11: ranija napomena „nema je u Brand Kit-u” nije tačna — fajl postoji u lokalnom Brand Kit-u (tmp/cosmos-lac-assets, grupa „uncertain” jer je isti fajl za više varijanti). Raniji tekst: '+base
        elif 'cosmoslac' not in base:
            base='Slika postoji samo na cosmoslac.com (nije u Brand Kit-u); upotreba traži pisanu dozvolu. '+base
    return re.sub(r'\s+',' ',base).strip() or 'Nema fotografije proizvoda na sajtu (prikazuje se „Vizuel u pripremi”).'

def local_of(x):
    if x['slug']=='baslac-30-s510-s-serija':
        return "public/products/baslac/baslac-30-s510-s-serija.webp postoji, ali prikazuje drugi proizvod (35-M1020) — ne koristi se"
    return "Nema (pretraženi public/, review-assets/, assets/, tmp/, _incoming/)"
SYSTEMS={'rm-agilis-x-treme':'AGILIS X-TREME','rm-onyx-hd-tropical':'ONYX HD TROPICAL'}
def need_of(x):
    if x['slug'] in SYSTEMS:
        return "Odluka o prikazu SISTEMA "+SYSTEMS[x['slug']]+" (nema jednu ambalažu): npr. bez fotografije, zvanična slika sistema/kolekcije ili konkretna komponenta — NE fotografisati jednu limenku"
    if x['brand']=='sata':
        reason=(unres.get(x['slug']) or {}).get('reason','')
        if reason.startswith('Vlasnik je ranije prihvatio placeholder'):
            return "Ništa obavezno — placeholder je ranije prihvaćen; slika samo uz saglasnost SATA"
        if reason.startswith('Ranije utvrđeno: SATA ne objavljuje sliku'):
            return "Vaša fotografija (SATA ne objavljuje sliku ovog artikla)"
        return "Potvrda prava: pisana saglasnost SATA (slika postoji na sata.com)"
    if x['brand']=='cosmos-lac': return "Potvrda prava: pisana dozvola Cosmos Lac ili dopuna Brand Kit-a"
    if x['slug']=='baslac-30-s510-s-serija': return "Potvrda identiteta (Line 30 Topcoat ili 30 CV), pa Vaša fotografija ambalaže"
    if x['slug']=='rm-diamont-bezbojni-lak': return "Potvrda koji je to DIAMONT lak (šifra), pa Vaša fotografija ambalaže"
    return "Vaša fotografija ambalaže, ako je artikal u našoj ponudi (proizvođač ne objavljuje sliku ove šifre)"
missing=[x for x in rows if 'placeholder' in (x['image'] or '') and not (x['brand']=='baslac' and x['redirect'])]
pending=[x for x in rows if 'pending-rights' in (x['image'] or '') or (x['brand']=='rupes' and 'placeholder' not in (x['image'] or ''))]
missing.sort(key=lambda x:(ORDER[x['brand']],x['name'].lower()))
pending.sort(key=lambda x:(ORDER[x['brand']],x['name'].lower()))
assert len(missing)==189, len(missing)
assert not set(x['slug'] for x in missing)&set(x['slug'] for x in pending), "preklapanje"

wb=Workbook()
HDR=Font(bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F2937"); EMPTY=PatternFill("solid",fgColor="FFF7D6")
def sheet(ws, headers, data, widths, empty_cols):
    ws.append(headers)
    for c in ws[1]: c.font=HDR; c.fill=FILL; c.alignment=Alignment(wrap_text=True,vertical="top")
    for row in data: ws.append(row)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    link_col=headers.index("Link naše stranice")+1
    for r in range(2, ws.max_row+1):
        for c in range(1,len(headers)+1):
            cell=ws.cell(r,c); cell.alignment=Alignment(wrap_text=True,vertical="top")
            if headers[c-1] in empty_cols: cell.fill=EMPTY
        lc=ws.cell(r,link_col)
        if lc.value: lc.hyperlink=lc.value; lc.font=Font(color="0563C1",underline="single")
    ws.freeze_panes="C2"; ws.auto_filter.ref=ws.dimensions
EMPTYC=["Pronađeni link","Naziv preuzetog fajla","Status"]
NAL=json.load(open('nalazi.json'))
NF={'rm':'Ne — pretraženo: Surventis Brand Portal R-M hub 51 (2026-10-06), info.rmpaint.com, rmpaint.com (sitemap, 303 stranice proizvoda, 2026-10-11)',
    'baslac':'Ne — pretraženo: Surventis Brand Portal baslac hub 54 (441 asset, 2026-10-06), baslac.com (stranica dodataka navodi šifru bez slike, 2026-10-11)',
    'norbin':'Ne — norbin-paint.com ne objavljuje slike; distributer Väritikka: N95-060 samo logo, ostale šifre ne postoje (2026-10-11)'}
def row_status(x):
    f=NAL.get(x['slug'])
    if f: return f
    found=NF.get(x['brand'],'Ne')
    if x['slug']=='rm-ra-040-agilis-mix': found+='; rmpaint.com „AGILIS MIX” prikazuje RA 050, ne RA 040 — odbačeno'
    if x['slug'] in ('rm-diamont-bezbojni-lak','baslac-30-s510-s-serija'):
        return dict(found=found,url='',page='',file='',sha='',match='ČEKA POTVRDU IDENTITETA',rights='—',need=need_of(x))
    return dict(found=found,url='',page='',file='',sha='',match='—',rights='—',need=None)
def final_need(x,f):
    if f.get('need'): return f['need']
    return need_of(x)
def status_of(x,f):
    if x['slug'] in SYSTEMS: return 'nije pronađeno — čeka odluku o prikazu sistema'
    if f['match'].startswith('ČEKA'): return 'čeka potvrdu identiteta'
    if not f['file']:
        n=final_need(x,f)
        return 'nije pronađeno — '+('placeholder prihvaćen' if n.startswith('Ništa') else 'potrebna Vaša fotografija (ako je u ponudi)')
    if not f['rights'].startswith('DA'):
        if f['match'].startswith('Potvrđeno'): return 'pronađeno — čeka pravo'
        if x['brand']=='cosmos-lac': return 'pronađeno — čeka pravo i Vašu odluku o zajedničkoj slici'
        return 'kandidat (slika srodnog artikla) — bez uvoza dok se ne potvrdi šta prikazuje; čeka i pravo'
    if f['match'].startswith('Potvrđeno'): return 'pronađeno — spremno za uvoz'
    return 'pronađeno — čeka Vašu odluku o zajedničkoj slici'
ws=wb.active; ws.title="Nedostaje fotografija"
PRIV=os.environ.get("PRIVATNO")=="1"
PROG=json.load(open("priv/program-detalji.json")) if PRIV else {}
H=["#","Brend","Tačan naziv","Šifra proizvođača","Varijanta / pakovanje","Link naše stranice"]+(["Aktuelno u našem programu","Pronađeno u pregledanom šifarniku (lager 23.07. / stanje 09.10.2026)","Trenutno stanje magacina","Istorijska evidencija (lager 23.07.2026)"] if PRIV else [])+["1. Fotografija pronađena i preuzeta","2. Podudaranje proizvod / varijanta / pakovanje","3. Pravo objavljivanja","Izvorni URL slike","Stranica izvora","Lokalni fajl (za pregled, NIJE objavljen)","SHA-256","Potrebno od Vas","Predloženi upit za pretragu","Napomena iz izveštaja"]+EMPTYC
data=[]
for i,x in enumerate(missing,1):
    code,cnote=code_of(x)
    note=note_of(x)+(f" ({cnote})" if cnote else "")
    f=row_status(x)
    st=status_of(x,f)
    data.append([i,BR[x['brand']],x['name'],code,pack_of(x),SITE+x['href']]+([PROG[x['slug']]['aktuelno'],PROG[x['slug']]['sifarnik_tekst'],PROG[x['slug']]['magacin_sada'],PROG[x['slug']]['istorija_jul']] if (PRIV and x['slug'] in PROG) else (['nije provereno (van spiska 25)','nije provereno','nepoznato','—'] if PRIV else []))+[f['found'],f['match'],f['rights'],f['url'],f['page'],
                 ('~/Carsystem/radno/slike-za-pronalazenje/'+f['file']) if f['file'] else '',f['sha'][:16],final_need(x,f),query_of(x,code),note,
                 f['url'] or '', os.path.basename(f['file']) if f['file'] else '', st])
sheet(ws,H,data,[5,10,40,16,18,38]+([30,46,22,26] if PRIV else [])+[40,48,40,40,34,46,18,36,34,50,34,30,30],EMPTYC)
STATUS=[r[-1] for r in data]
json.dump({x['slug']:{'status':st,'need':r[H.index('Potrebno od Vas')]} for x,st,r in zip(missing,STATUS,data)},open('statusi.json','w'),ensure_ascii=False,indent=1)

ws2=wb.create_sheet("Čeka potvrdu prava")
H2=["#","Brend","Tačan naziv","Šifra proizvođača","Varijanta / pakovanje","Link naše stranice","Slika koju već imamo (fajl u repou)","Poreklo slike","Šta treba potvrditi","Gde se danas prikazuje"]+["Odluka (da/ne)","Status"]
data2=[]
for i,x in enumerate(pending,1):
    code,cnote=code_of(x)
    if x['brand']=='norbin':
        f="review-assets/pending-rights/norbin/"+x['image'].rsplit('/',1)[-1]
        src="Render proizvođača sa sajta distributera Väritikka"
        what="Pravo da render objavimo na sajtu (rightsBasis OWNER_CONFIRMATION_REQUIRED)"
        where="Lokalno i na Preview-u da; na Production-u NE (prikazuje se „Vizuel u pripremi”)"
    else:
        f="public"+x['image']; src="carsystem.org (Carsystem katalog), proizvod pod RUPES brendom"
        what="Pravo prikaza RUPES proizvoda slikom sa carsystem.org (RUPES_IMAGE_RIGHTS_REVIEW)"
        where="Prikazuje se svuda (i na Production-u)"
    data2.append([i,BR[x['brand']],x['name'],code,pack_of(x),SITE+x['href'],f,src+(f" — {cnote}" if cnote else ""),what,where,"",""])
sheet(ws2,H2,data2,[5,10,42,18,22,40,48,40,48,36,14,14],["Odluka (da/ne)","Status"])

# --- ambalaža po zapremini (baslac sistemske strane) ---
INV=json.load(open('audit/inventory.json'))
fam=lambda p:p['canonical'].split('?')[0].split('/')[-1]
SYS={'baslac-line-30':'Line 30','baslac-line-35':'Line 35','baslac-line-45':'Line 45'}
bases=[p for p in INV['products'] if p['brand']=='baslac' and p['redirect'] and fam(p) in SYS]
SLOT={('Line 30','1 L'):'F-30-1L',('Line 35','1 L'):'F-35-1L',('Line 35','0,5 L'):'F-35-05L',('Line 45','5 L'):'F-45-5L',('Line 45','1 L'):'F-45-1L',('Line 45','0,5 L'):'F-45-05L',('Line 45','0,1 L'):'F-45-01L'}
rowsB=[]
for p in sorted(bases,key=lambda p:(fam(p),p['name'])):
    vol='/'.join(p['packages']) or (p['volume'] or '')
    key=(SYS[fam(p)],vol); slot=SLOT.get(key)
    if not slot: continue
    ph='placeholder' in (p['image'] or '')
    rowsB.append([SYS[fam(p)],vol,slot,p['sku'],p['name'],SITE+p['canonical'],'da' if ph else 'ne (prikazuje se primer limenke linije 45-W)','DA' if ph else 'ne'])
cntV=collections.Counter(r[2] for r in rowsB); cnt66=collections.Counter(r[2] for r in rowsB if r[7]=='DA')
ws4=wb.create_sheet("baslac ambalaža po zapremini")
ws4.append(["Slotovi: BROJ VARIJANTI ≠ BROJ POTREBNIH FOTOGRAFIJA. Svaka fotografija je jedna generička limenka sistema u datoj zapremini; baze se razlikuju šifrom i nijansom, ne limenkom."])
ws4.append([])
ws4.append(["Slot fotografije","Sistem","Pakovanje","Varijanti (baza) u slotu","Od toga u spisku od 66 (bez ikakve slike)","Potrebnih različitih fotografija","Postojeći materijal","Potrebno od Vas"])
NOTE={'F-30-1L':'Na Surventis portalu postoji sopstveni 1 L packshot SAMO za bazu 30-S110 (asset 12114) — pokriva tu jednu varijantu, ne slot','F-45-1L':'Generička „45-W” limenka koju je vlasnik dostavio (prikazuje se kao primer)'}
for slot in SLOT.values():
    sy,pk=[k for k,v in SLOT.items() if v==slot][0]
    ws4.append([slot,sy,pk,cntV.get(slot,0),cnt66.get(slot,0),1,NOTE.get(slot,'Generička limenka ove zapremine ne postoji na portalu' if sy!='Line 45' else 'Prikazuje se primer limenke 45-W (druga zapremina); tačna zapremina nedostaje'),f'Jedna fotografija prazne/generičke limenke {sy} {pk}'])
ws4.append(["UKUPNO","","",sum(cntV.values()),sum(cnt66.values()),len(SLOT),"",""])
ws4.append([]); ws4.append(["Mapiranje varijanti na slotove (sve baze sistemskih strana; kolona „u spisku 66” označava redove iz spiska druge sesije)"])
hdr=["Sistem","Pakovanje","Slot fotografije","Šifra baze","Naziv","Link (varijanta)","Zapis bez ikakve slike","U spisku 66"]
ws4.append(hdr)
hr=ws4.max_row
for c in ws4[3]+ws4[hr]: c.font=HDR; c.fill=FILL; c.alignment=Alignment(wrap_text=True,vertical="top")
for r in rowsB: ws4.append(r)
for i,w in enumerate([12,10,12,14,40,60,26,10],1): ws4.column_dimensions[get_column_letter(i)].width=w
for row in ws4.iter_rows(min_row=3):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
ws4['A1'].font=Font(bold=True)
BASLAC_SUM=(sum(cntV.values()),sum(cnt66.values()),len(SLOT),dict(cnt66))
# --- pregled ---
import collections
ws3=wb.create_sheet("Pregled",0)
cnt=collections.Counter(STATUS)
byb=collections.defaultdict(collections.Counter)
for x,st in zip(missing,STATUS): byb[BR[x['brand']]][st]+=1
lines=[["Pregled pretrage fotografija — 2026-10-11"],[""],
 ["Glavni spisak","189 proizvoda bez fotografije (list „Nedostaje fotografija”) + 7 slotova ambalaže baslac sistema (poseban list) + 31 slika koja čeka prava (list „Čeka potvrdu prava”). Ovo je JEDINI glavni spisak; zamenjuje radni fajl ~/Carsystem/radno/slike-za-pronalazenje/SLIKE_ZA_PRONALAZENJE.xlsx (255 redova) iz druge sesije."],
 ["66 baslac redova iz druge sesije","To su baze sa sistemskih strana Line 30 (17 × 1 L) i Line 35 (17 × 1 L + 32 × 0,5 L). Nisu posebni proizvodi bez slike: baze se kupcu prikazuju na sistemskoj strani, a nedostaje jedna fotografija generičke limenke po pakovanju. Isti nalaz je već evidentiran kao 3 slota ambalaže (OLD-043…045): 66 varijanti = 3 potrebne fotografije (Line 30 1 L: 17, Line 35 1 L: 17, Line 35 0,5 L: 32). Sa Line 45 (68 baza, prikazuje se primer limenke) ukupno 134 baze = 7 fotografija. Mapiranje svake varijante je na listu „baslac ambalaža po zapremini”. Nisu ponovljeni nalazi među 189 — preklapanje je 0."],
 ["FO-314","Prebačeno u „čeka potvrdu identiteta”: etiketa ne nosi šifru ni nijansu, Brand Kit isti fajl nosi i kao FO-313 (zvanični FO-313 je Cherry Dark), a poklopac je crven, bliži FO-312/313 nego pink nijansama 308–310."],
 ["Cosmos Brand Kit — pravo","ISPRAVKA: ranije označeno kao potvrđeno. Brand Kit je zvaničan izvor (742 objavljene slike), ali pisana dozvola za web upotrebu nije dokumentovana — vidi nacrt zahteva Cosmos Lac-u."],
 ["Ponuda i magacin","Za 25 proizvoda bez pronađene slike ponuda nije potvrđena; odsustvo iz pregledanog šifarnika NE znači uklanjanje sa sajta. Trenutno stanje magacina je nepoznato dok ga vlasnik ne proveri (spisak: MAGACIN_PROVERA.xlsx, privatno)."],
 ["Sistemi bez šifre","AGILIS X-TREME i ONYX HD TROPICAL: nema zahteva za fotografiju jedne ambalaže — traži se odluka o prikazu sistema."],
 ["SATA — slike srodnih artikala","7 slika su samo kandidati: ne uvoze se dok se ne potvrdi šta tačno prikazuju (i dok nema prava)."],
 ["Pregled sa slikama","~/Carsystem/radno/slike-za-pronalazenje/PREGLED_SPORNIH_SLIKA.html (i .jpg): 8 Cosmos sa zajedničkom slikom, 7 SATA sa slikom srodnog artikla, FO-314; neslaganje 175/185 ml za Chrome 450 / Gold 451 je istaknuto — katalog NIJE menjan."],
 [""],["Stanje","Broj"]]+[[k,v] for k,v in sorted(cnt.items(),key=lambda t:-t[1])]+[[""],["Po brendu","Stanja"]]+[[b, "; ".join(f"{k}: {v}" for k,v in c.items())] for b,c in byb.items()]+[[""],
 ["Gde su fajlovi","~/Carsystem/radno/slike-za-pronalazenje/preuzeto/<brend>/ — van git-a, nije objavljeno. SATA _manifest.json ima URL, SHA-256 i stranicu artikla za svaki fajl."],
 ["Pravila","Podudaranje „Potvrđeno” samo kad slika nosi tačnu šifru/naziv ili je sopstvena glavna slika artikla na stranici proizvođača. Zajednička slika porodice je označena posebno. Šifre se ne izmišljaju."]]
for r in lines: ws3.append(r)
ws3.column_dimensions['A'].width=34; ws3.column_dimensions['B'].width=130
for row in ws3.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    row[0].font=Font(bold=True)
wb.save(OUT); print("ok",len(missing),len(pending)); print(cnt); print({b:dict(c) for b,c in byb.items()})
