"""Sintetički BizniSoft izveštaj „STANJE ZALIHA" (XLSX + CSV) istog rasporeda kao izvoz
štampe (dok. 65): list Page1, zaglavlje izveštaja u redovima 1–14, zaglavlje tabele u redu 16
koje se ponavlja na svakoj strani, podnožje „Strana N od M", spojene ćelije (naslov „Količina"
u X, vrednost u Y), zbirni red „Ukupno stanje zaliha na dan:". Svi podaci su izmišljeni.

    python3 scripts/fixtures/build-stock-report-fixture.py fixtures/dev/biznisoft/lager
"""
import csv, os, sys

sys.path.insert(0, os.path.expanduser("~/Library/Python/3.9/lib/python/site-packages"))
import openpyxl  # noqa: E402
from openpyxl.worksheet.pagebreak import Break  # noqa: E402

OUT = sys.argv[1]
ROWS = [
    # (šifra, naziv, količina, nabavna, VP)
    (900001, "SINT BAZA SREBRNA 1 L", 12, 1000, 1300),
    (900002, "SINT LAK HS 5 L", 7.5, 4000, 5200),
    (900003, "SINT HERTER 2,5 L", 2.5, 1500, 1950),
    (900004, "SINT KIT 1 KG", 0, 800, 1000),
    (900005, "SINT BRUSNI DISK P120", 230, 20, 30),
    (900006, "SINT TRAKA 19 MM", 48, 150, 210),
    (12, "SINT STARA ŠIFRA BEZ NULA", 3, 100, 130),
    (900008, "SINT ČISTAČ ĐAK ŽUTI", 1.25, 600, 780),
    (900009, "SINT RUKAVICE NITRIL", 4, 900, 1200),
    (900010, "SINT PIŠTOLJ", 1, 20000, 26000),
]
PER_PAGE = 4

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Page1"
ws["A1"] = "CAR SYSTEM SINT"
ws["AP1"] = "Datum štampe: 05.09.2026"
ws["M2"] = "STANJE ZALIHA"
ws["A4"] = "SINT MESTO"
ws["B7"], ws["H7"], ws["O7"], ws["V7"] = "Na dan:", "05.09.2026", "Način sorta:", "Po nazivu"
ws["B9"], ws["J9"] = "Vrsta artikla: ", "Trgovačka roba, Proizvodi, "
ws["B10"], ws["K10"], ws["T10"], ws["AE10"] = "Vrsta nabavne cene:", "Prosečna nabavna cena", "Stanje artikla:", "Artikli na zalihama"
ws["A12"] = "Poslovni objekti:"
ws["A14"] = "021 VELEPRODAJA"


def header(r):
    for col, txt in (("A", "Šifra"), ("F", "Naziv robe"), ("R", "Stopa"), ("X", "Količina"), ("AD", "Nab.Cena"), ("AI", "Nab.Vred"), ("AL", "VP Cena"), ("AT", "VP Vred."), ("AV", "%RuC")):
        ws[f"{col}{r}"] = txt
    ws.merge_cells(f"A{r}:E{r}")
    ws.merge_cells(f"F{r}:Q{r}")


r = 16
header(r)
r += 2
pages = (len(ROWS) + PER_PAGE - 1) // PER_PAGE
sum_nab = sum_vp = 0
for i, (code, name, qty, nab, vp) in enumerate(ROWS):
    if i and i % PER_PAGE == 0:
        ws[f"AJ{r}"] = "BizniSoft Poslovni Programi - www.biznisoft.com"
        ws[f"A{r + 1}"] = f"Strana {i // PER_PAGE} od {pages}"
        ws.row_breaks.append(Break(id=r + 1))
        r += 4
        header(r)
        r += 2
    ws[f"A{r}"], ws[f"F{r}"], ws[f"R{r}"], ws[f"Y{r}"] = code, name, "20 %", qty
    ws[f"AD{r}"], ws[f"AI{r}"], ws[f"AL{r}"], ws[f"AT{r}"] = nab, round(qty * nab, 2), vp, round(qty * vp, 2)
    ws[f"AV{r}"] = round((vp - nab) / nab * 100, 2)
    ws.merge_cells(f"F{r}:Q{r}")
    ws.merge_cells(f"Y{r}:AC{r}")
    sum_nab += qty * nab
    sum_vp += qty * vp
    r += 1
ws[f"C{r + 1}"], ws[f"AA{r + 1}"], ws[f"AM{r + 1}"] = "Ukupno stanje zaliha na dan:", f"{sum_nab:.2f}", f"{sum_vp:.2f}"
ws[f"AJ{r + 3}"] = "BizniSoft Poslovni Programi - www.biznisoft.com"
ws[f"A{r + 4}"] = f"Strana {pages} od {pages}"
os.makedirs(OUT, exist_ok=True)
wb.save(os.path.join(OUT, "stanje-zaliha-sinteticki.xlsx"))

# CSV istog sadržaja: decimalni zarez, separator ';', Windows-1250 (kako piše BizniSoft na Windows-u).
with open(os.path.join(OUT, "stanje-zaliha-sinteticki.csv"), "w", encoding="cp1250", newline="") as f:
    w = csv.writer(f, delimiter=";")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
        out = []
        for v in row:
            if v is None:
                out.append("")
            elif isinstance(v, float):
                out.append(f"{v:.2f}".replace(".", ","))
            elif isinstance(v, str) and v.replace(".", "", 1).isdigit() and "." in v:
                out.append(v.replace(".", ","))
            else:
                out.append(str(v))
        w.writerow(out)
print("ok", ws.max_row, "redova")
