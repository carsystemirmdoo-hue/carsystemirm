import json, csv, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
REPO="/Users/miledulic/Carsystem/radno/javni-sajt-ispravke"
OUT=sys.argv[1]
rows=json.load(open('slike.json'))
unres={}
for r in csv.DictReader(open(f"{REPO}/docs/catalog/image-gaps-2026-10/UNRESOLVED_IMAGES.csv")):
    slug=r['image_id'].split('__',1)[-1] if r['image_id'] else ''
    unres[slug]=r
sata={r['slug']:r for r in csv.DictReader(open(f"{REPO}/docs/catalog/image-gaps-2026-10/sata-current-official-urls.csv"))}
BR={"rm":"R-M","baslac":"baslac","norbin":"Norbin","cosmos-lac":"Cosmos Lac","sata":"SATA","rupes":"RUPES"}
ORDER={"rm":0,"baslac":1,"norbin":2,"cosmos-lac":3,"sata":4,"rupes":5}
SITE="https://carsystemirm.com"
INTERNAL=re.compile(r'^(RM-|NORBIN-|CL-|SATA-)')
def code_of(x):
    code,note=_code_of(x)
    if code and not re.search(r'\d', code):
        return "", "sistem bez šifre proizvođača"
    if code and x['brand']=='rupes' and re.fullmatch(r'\d{3}\.\d{3}', code):
        return "", f"Carsystem broj artikla {code}; RUPES šifra nije potvrđena"
    return code,note
def _code_of(x):
    if x['brand']=='sata' and x['slug'] in sata and sata[x['slug']]['main_article']:
        return sata[x['slug']]['main_article'], "SATA broj artikla (sa sata.com)"
    if x['manufacturerCode']: return x['manufacturerCode'], ""
    for s in x['specs']:
        lab,val=s.split(': ',1)
        if val and not INTERNAL.match(val): return val, ""
    if x['publicCode'] and not INTERNAL.match(x['publicCode']): return x['publicCode'], ""
    if x['externalSku'] and not INTERNAL.match(x['externalSku']): return x['externalSku'], ""
    return "", "šifra proizvođača nije potvrđena"
def pack_of(x):
    parts=[]
    if x['variantOptions']: parts+= [v for v in x['variantOptions'].values() if v]
    p=[l for l in x['packages'] if l and l!='Na upit']
    parts+=[l for l in p if l not in parts]
    if x['volume'] and x['volume'] not in parts: parts.append(x['volume'])
    return ", ".join(parts)
def strip_brand(name,b):
    return re.sub(r'^(R-M|baslac|Baslac|Norbin|Cosmos Lac|SATA|RUPES)\s+','',name)
def query_of(x,code):
    b=BR[x['brand']]; n=strip_brand(x['name'],b)
    if x['brand']=='sata': return f'SATA {code} {n}'.strip()
    q=f'{b} {code} {n}' if code and code not in n else f'{b} {n}'
    return re.sub(r'\s+',' ',q).strip()+' packshot'
def source_of(x):
    if x['brand']=='sata' and x['slug'] in sata:
        s=sata[x['slug']]; return f"Stranica artikla: {s['product_page_url']} · zvanična slika: {s['current_image_url']}"
    u=unres.get(x['slug'])
    return (u['checked_sources'] if u else '')
def note_of(x):
    u=unres.get(x['slug'])
    base=u['reason'] if u else ''
    if x['brand']=='sata':
        if u and 'PLACEHOLDER_ACCEPTED' in (u['state']+u['reason']+u.get('group','')): base='Placeholder ranije prihvaćen (pribor). '+base
        pass
    if x['brand']=='cosmos-lac': base='Slika postoji samo na cosmoslac.com (nije u Brand Kit-u); upotreba traži pisanu dozvolu. '+base if 'cosmoslac' not in base else base
    return re.sub(r'\s+',' ',base).strip() or 'Nema fotografije proizvoda na sajtu (prikazuje se „Vizuel u pripremi”).'

def local_of(x):
    if x['slug']=='baslac-30-s510-s-serija':
        return "public/products/baslac/baslac-30-s510-s-serija.webp postoji, ali prikazuje drugi proizvod (35-M1020) — ne koristi se"
    return "Nema (pretraženi public/, review-assets/, assets/, tmp/, _incoming/)"
def need_of(x):
    if x['brand']=='sata':
        reason=(unres.get(x['slug']) or {}).get('reason','')
        if reason.startswith('Vlasnik je ranije prihvatio placeholder'):
            return "Ništa obavezno — placeholder je ranije prihvaćen; slika samo uz saglasnost SATA"
        if reason.startswith('Ranije utvrđeno: SATA ne objavljuje sliku'):
            return "Vaša fotografija (SATA ne objavljuje sliku ovog artikla)"
        return "Potvrda prava: pisana saglasnost SATA (slika postoji na sata.com)"
    if x['brand']=='cosmos-lac': return "Potvrda prava: pisana dozvola Cosmos Lac ili dopuna Brand Kit-a"
    if x['slug']=='baslac-30-s510-s-serija': return "Potvrda identiteta (Line 30 Topcoat ili 30 CV), pa Vaša fotografija ambalaže"
    if x['slug']=='rm-diamont-bezbojni-lak': return "Potvrda koji je to DIAMONT lak (šifra), pa Vaša fotografija ambalaže"
    return "Vaša fotografija ambalaže (proizvođač ne objavljuje sliku ove šifre)"
missing=[x for x in rows if 'placeholder' in (x['image'] or '') and not (x['brand']=='baslac' and x['redirect'])]
pending=[x for x in rows if 'pending-rights' in (x['image'] or '') or (x['brand']=='rupes' and 'placeholder' not in (x['image'] or ''))]
missing.sort(key=lambda x:(ORDER[x['brand']],x['name'].lower()))
pending.sort(key=lambda x:(ORDER[x['brand']],x['name'].lower()))
assert len(missing)==189, len(missing)
assert not set(x['slug'] for x in missing)&set(x['slug'] for x in pending), "preklapanje"

wb=Workbook()
HDR=Font(bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F2937"); EMPTY=PatternFill("solid",fgColor="FFF7D6")
def sheet(ws, headers, data, widths, empty_cols):
    ws.append(headers)
    for c in ws[1]: c.font=HDR; c.fill=FILL; c.alignment=Alignment(wrap_text=True,vertical="top")
    for row in data: ws.append(row)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    link_col=headers.index("Link naše stranice")+1
    for r in range(2, ws.max_row+1):
        for c in range(1,len(headers)+1):
            cell=ws.cell(r,c); cell.alignment=Alignment(wrap_text=True,vertical="top")
            if headers[c-1] in empty_cols: cell.fill=EMPTY
        lc=ws.cell(r,link_col)
        if lc.value: lc.hyperlink=lc.value; lc.font=Font(color="0563C1",underline="single")
    ws.freeze_panes="C2"; ws.auto_filter.ref=ws.dimensions
EMPTYC=["Pronađeni link","Naziv preuzetog fajla","Status"]
ws=wb.active; ws.title="Nedostaje fotografija"
H=["#","Brend","Tačan naziv","Šifra proizvođača","Varijanta / pakovanje","Link naše stranice","Predloženi upit za pretragu","Šta nedostaje (napomena)","Poznat zvanični izvor (ranija provera)","Postojeći lokalni materijal","Potrebno od Vas"]+EMPTYC
data=[]
for i,x in enumerate(missing,1):
    code,cnote=code_of(x)
    note=note_of(x)+(f" ({cnote})" if cnote else "")
    data.append([i,BR[x['brand']],x['name'],code,pack_of(x),SITE+x['href'],query_of(x,code),note,source_of(x),local_of(x),need_of(x),"","",""])
sheet(ws,H,data,[5,11,42,18,22,40,40,60,50,34,34,32,26,14],EMPTYC)

ws2=wb.create_sheet("Čeka potvrdu prava")
H2=["#","Brend","Tačan naziv","Šifra proizvođača","Varijanta / pakovanje","Link naše stranice","Slika koju već imamo (fajl u repou)","Poreklo slike","Šta treba potvrditi","Gde se danas prikazuje"]+["Odluka (da/ne)","Status"]
data2=[]
for i,x in enumerate(pending,1):
    code,cnote=code_of(x)
    if x['brand']=='norbin':
        f="review-assets/pending-rights/norbin/"+x['image'].rsplit('/',1)[-1]
        src="Render proizvođača sa sajta distributera Väritikka"
        what="Pravo da render objavimo na sajtu (rightsBasis OWNER_CONFIRMATION_REQUIRED)"
        where="Lokalno i na Preview-u da; na Production-u NE (prikazuje se „Vizuel u pripremi”)"
    else:
        f="public"+x['image']; src="carsystem.org (Carsystem katalog), proizvod pod RUPES brendom"
        what="Pravo prikaza RUPES proizvoda slikom sa carsystem.org (RUPES_IMAGE_RIGHTS_REVIEW)"
        where="Prikazuje se svuda (i na Production-u)"
    data2.append([i,BR[x['brand']],x['name'],code,pack_of(x),SITE+x['href'],f,src+(f" — {cnote}" if cnote else ""),what,where,"",""])
sheet(ws2,H2,data2,[5,10,42,18,22,40,48,40,48,36,14,14],["Odluka (da/ne)","Status"])

ws3=wb.create_sheet("Napomene")
notes=[["Izvor","JAVNI_SAJT_NEDOSTACI.csv / JAVNI_SAJT_INVENTAR_PROIZVODA.csv (provera 2026-10-11), podaci iz trenutnog koda grane fix/javni-sajt-2026-10-11"],
["List „Nedostaje fotografija”",f"{len(missing)} proizvoda bez fotografije na sajtu: "+", ".join(f"{BR[b]} {n}" for b,n in sorted(__import__('collections').Counter(x['brand'] for x in missing).items(),key=lambda t:ORDER[t[0]]))],
["List „Čeka potvrdu prava”",f"{len(pending)} slika koje već imamo: Norbin {sum(1 for x in pending if x['brand']=='norbin')} (na Production-u se kriju) i RUPES {sum(1 for x in pending if x['brand']=='rupes')} (prikazuju se). Ovih 6 Norbin proizvoda NIJE među {len(missing)} — preklapanje je provereno (0)."],
["Šifre","Šifra je upisana samo kada je potvrđena u katalogu (oznaka proizvođača ili SATA broj artikla). Prazno = nije potvrđena; ne pogađati."],
["Nije uključeno","baslac sistemske strane (Line 30/35/45): 156 baza ima generičku limenku linije; nedostaje fotografija prazne limenke po zapremini (1 L, 0,5 L, 0,1 L, 5 L) — to nije fotografija proizvoda, vidi izveštaj (G02)."],
["Žuta polja","Pronađeni link, Naziv preuzetog fajla, Status — popunjavate Vi."]]
for r in notes: ws3.append(r)
ws3.column_dimensions['A'].width=26; ws3.column_dimensions['B'].width=120
for row in ws3.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    row[0].font=Font(bold=True)
wb.save(OUT); print("ok",len(missing),len(pending))
